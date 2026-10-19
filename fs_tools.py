"""Local filesystem tools for reading, listing, writing, and searching resumes."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile


def _file_metadata(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "name": path.name,
        "size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(
            stat.st_mtime, tz=timezone.utc
        ).isoformat(),
        "extension": path.suffix.lower(),
    }


def _extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".txt":
        return path.read_text(encoding="utf-8-sig")

    if extension == ".pdf":
        try:
            from pypdf import PdfReader
            from pypdf.errors import FileNotDecryptedError, PdfReadError
        except ImportError as exc:
            raise RuntimeError("PDF support requires the 'pypdf' package.") from exc
        try:
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except (FileNotDecryptedError, PdfReadError) as exc:
            raise ValueError(f"Unable to extract PDF text: {exc}") from exc

    if extension == ".docx":
        try:
            from docx import Document
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as exc:
            raise RuntimeError(
                "DOCX support requires the 'python-docx' package."
            ) from exc
        try:
            document = Document(str(path))
        except (BadZipFile, PackageNotFoundError) as exc:
            raise ValueError(f"Unable to read DOCX document: {exc}") from exc
        content = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            content.extend(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        return "\n".join(line for line in content if line)

    raise ValueError(
        f"Unsupported file type '{path.suffix or '(no extension)'}'. "
        "Supported types are PDF, TXT, and DOCX."
    )


def read_file(filepath: str) -> dict[str, Any]:
    """Extract text and metadata from a PDF, TXT, or DOCX file."""
    path = Path(filepath).expanduser()
    try:
        if not path.is_file():
            return {
                "success": False,
                "filepath": str(path),
                "error": "File does not exist or is not a regular file.",
            }
        content = _extract_text(path)
        return {
            "success": True,
            "filepath": str(path),
            "content": content,
            "metadata": _file_metadata(path),
        }
    except (OSError, UnicodeError, ValueError, RuntimeError) as exc:
        return {"success": False, "filepath": str(path), "error": str(exc)}


def list_files(directory: str, extension: str | None = None) -> list[dict[str, Any]]:
    """List direct child files in a directory, optionally filtered by extension."""
    path = Path(directory).expanduser()
    if extension:
        suffix = extension.lower()
        if not suffix.startswith("."):
            suffix = f".{suffix}"
    else:
        suffix = None

    try:
        if not path.is_dir():
            return [{"error": f"Directory does not exist: {path}"}]
        results = []
        for item in sorted(path.iterdir(), key=lambda entry: entry.name.lower()):
            if not item.is_file() or (suffix and item.suffix.lower() != suffix):
                continue
            try:
                results.append({"filepath": str(item), **_file_metadata(item)})
            except OSError as exc:
                results.append({"filepath": str(item), "error": str(exc)})
        return results
    except OSError as exc:
        return [{"error": f"Unable to list directory '{path}': {exc}"}]


def write_file(filepath: str, content: str) -> dict[str, Any]:
    """Write UTF-8 text to a file, creating missing parent directories."""
    path = Path(filepath).expanduser()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return {
            "success": True,
            "filepath": str(path),
            "bytes_written": path.stat().st_size,
        }
    except (OSError, UnicodeError) as exc:
        return {"success": False, "filepath": str(path), "error": str(exc)}


def search_in_file(filepath: str, keyword: str) -> dict[str, Any]:
    """Find case-insensitive occurrences and return a short excerpt for each."""
    if not keyword:
        return {"success": False, "filepath": filepath, "error": "Keyword is empty."}

    result = read_file(filepath)
    if not result["success"]:
        return result

    content = result["content"]
    matches = []
    for match in re.finditer(re.escape(keyword), content, flags=re.IGNORECASE):
        start = max(0, match.start() - 60)
        end = min(len(content), match.end() + 60)
        matches.append(
            {
                "match": match.group(0),
                "position": match.start(),
                "context": content[start:end].replace("\n", " ").strip(),
            }
        )

    return {
        "success": True,
        "filepath": result["filepath"],
        "keyword": keyword,
        "match_count": len(matches),
        "matches": matches,
        "metadata": result["metadata"],
    }
